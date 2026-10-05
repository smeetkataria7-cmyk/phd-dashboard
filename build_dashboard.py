"""Build a self-contained HTML dashboard from the first N sheets of an Excel workbook.

Usage:
    python build_dashboard.py "PHD DATASET.xlsx"                 # -> dashboard.html next to the script
    python build_dashboard.py "NEW DATASET.xlsx" -o out.html     # custom output
    python build_dashboard.py "NEW DATASET.xlsx" -n 7            # number of leading sheets to use (default 7)

Works on any workbook (generic profiling + chart builder per sheet). When the sugarcane trial
schema is recognised (Brix/Pol/CCS, NDVI, weather, field lat/long ...) the Overview adds
domain panels automatically; panels whose columns are missing are simply skipped.
"""
import argparse
import datetime as dt
import json
import math
import re
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}")


# --------------------------------------------------------------------------- loading
def read_sheet_rows(ws, max_blank_run=300):
    """Read rows until a long run of fully blank rows (some sheets have ~1M formatted-but-empty rows)."""
    rows, blank_run, last_filled = [], 0, 0
    for r in ws.iter_rows(values_only=True):
        if any(v is not None and str(v).strip() != "" for v in r):
            blank_run = 0
            last_filled = len(rows) + 1
        else:
            blank_run += 1
        rows.append(r)
        if blank_run >= max_blank_run:
            break
    return rows[:last_filled]


def clean_name(v):
    return re.sub(r"\s+", " ", str(v)).strip()


def find_header(rows):
    """Header = first of the leading rows holding (almost) the maximum number of text cells."""
    scan = rows[:15]
    counts = [sum(isinstance(v, str) and v.strip() != "" for v in r) for r in scan]
    top = max(counts) if counts else 0
    for i, c in enumerate(counts):
        if c >= 0.8 * top and c > 0:
            return i
    return 0


def is_number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool) and not (isinstance(v, float) and math.isnan(v))


def to_float(v):
    if is_number(v):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.strip().replace(",", ""))
        except ValueError:
            return None
    return None


def iso(v):
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, dt.date):
        return v.isoformat()
    return None


def num_clean(x, nd=5):
    if x is None:
        return None
    if x == int(x) and abs(x) < 1e12:
        return int(x)
    return round(x, nd)


def classify(values):
    """Return (type, converted_values). type in numeric | date | cat | text."""
    present = [v for v in values if v is not None and not (isinstance(v, str) and v.strip() == "")]
    if not present:
        return "empty", [None] * len(values)
    n = len(present)
    n_date = sum(isinstance(v, (dt.datetime, dt.date)) for v in present)
    n_iso = sum(isinstance(v, str) and ISO_DATE.match(v.strip()) is not None for v in present)
    if n_date + n_iso >= 0.8 * n:
        out = []
        for v in values:
            if isinstance(v, (dt.datetime, dt.date)):
                out.append(iso(v))
            elif isinstance(v, str) and ISO_DATE.match(v.strip()):
                out.append(v.strip()[:10])
            else:
                out.append(None)
        return "date", out
    n_num = sum(to_float(v) is not None for v in present)
    return_num = n_num >= 0.9 * n
    return_vals = None
    if return_num:
        return_vals = [to_float(v) for v in values]
    if return_num:
        return "numeric", return_vals
    strs = [None if (v is None or (isinstance(v, str) and v.strip() == "")) else str(v).strip()[:80] for v in values]
    nuniq = len({s for s in strs if s is not None})
    return ("cat" if nuniq <= 30 else "text"), strs


def load_workbook_sheets(path, n_sheets):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    all_names = wb.sheetnames
    sheets = []
    for name in all_names[:n_sheets]:
        rows = read_sheet_rows(wb[name])
        if not rows:
            continue
        h = find_header(rows)
        header = [clean_name(c) if c is not None else "" for c in rows[h]]
        body = rows[h + 1:]
        ncols = len(header)
        # drop sparse rows (units / description rows) and blank rows
        min_filled = max(2, int(0.15 * sum(1 for x in header if x)))
        body = [r for r in body if sum(v is not None and str(v).strip() != "" for v in r[:ncols]) >= min_filled]
        # build columns
        seen, cols, data = {}, [], {}
        for j, nm in enumerate(header):
            if not nm:
                continue
            if nm in seen:
                seen[nm] += 1
                nm = f"{nm}.{seen[nm]}"
            else:
                seen[nm] = 0
            raw = [(r[j] if j < len(r) else None) for r in body]
            typ, vals = classify(raw)
            if typ == "empty":
                continue
            if typ == "numeric" and re.search(r"year", nm, re.I):
                typ, vals = "cat", [None if v is None else (str(int(v)) if v == int(v) else str(v)) for v in vals]
            data[nm] = vals
            cols.append({"name": nm, "type": typ})
        sheets.append({"name": name, "nrows": len(body), "cols": cols, "data": data})
    return all_names, sheets


# --------------------------------------------------------------------------- enrichment
VAR_PAREN = re.compile(r"\s*\(.*\)\s*$")


def enrich(sheets):
    """Normalise 'Variety Number' and derive it for sheets that only carry Plot_ID."""
    for s in sheets:
        d = s["data"]
        if "Variety Number" in d:
            full = [v for v in d["Variety Number"]]
            d["Variety (full)"] = full
            d["Variety Number"] = [None if v is None else VAR_PAREN.sub("", v).strip() for v in full]
            for c in s["cols"]:
                if c["name"] == "Variety Number":
                    c["type"] = "cat"
            if not any(c["name"] == "Variety (full)" for c in s["cols"]):
                s["cols"].append({"name": "Variety (full)", "type": "cat"})
    plot_map = {}
    for s in sheets:
        d = s["data"]
        if "Plot_ID" in d and "Variety Number" in d:
            for p, v in zip(d["Plot_ID"], d["Variety Number"]):
                if p and v:
                    plot_map.setdefault(p, v)
    for s in sheets:
        d = s["data"]
        if "Plot_ID" in d and "Variety Number" not in d and plot_map:
            d["Variety Number"] = [plot_map.get(p) for p in d["Plot_ID"]]
            s["cols"].append({"name": "Variety Number", "type": "cat", "derived": True})
    return sheets


def profile(sheets):
    for s in sheets:
        n = s["nrows"]
        for c in s["cols"]:
            vals = s["data"][c["name"]]
            present = [v for v in vals if v is not None]
            c["filled"] = round(len(present) / n, 4) if n else 0
            c["unique"] = len(set(present))
            if c["type"] == "numeric" and present:
                srt = sorted(present)
                c["min"], c["max"] = num_clean(srt[0]), num_clean(srt[-1])
                c["mean"] = round(sum(present) / len(present), 4)
                c["median"] = num_clean(srt[len(srt) // 2])
            elif c["type"] == "date" and present:
                c["min"], c["max"] = min(present), max(present)
            elif c["type"] == "cat":
                counts = {}
                for v in present:
                    counts[v] = counts.get(v, 0) + 1
                c["cats"] = sorted(counts.items(), key=lambda kv: -kv[1])[:30]
        # compact numeric payload
        for c in s["cols"]:
            if c["type"] == "numeric":
                s["data"][c["name"]] = [num_clean(v) for v in s["data"][c["name"]]]
    return sheets


# --------------------------------------------------------------------------- html
def build(path, n_sheets, out):
    all_names, sheets = load_workbook_sheets(path, n_sheets)
    sheets = profile(enrich(sheets))
    payload = {
        "file": Path(path).name,
        "total_sheets": len(all_names),
        "used_sheets": [s["name"] for s in sheets],
        "skipped_sheets": all_names[n_sheets:],
        "sheets": sheets,
    }
    html = (HERE / "dashboard_template.html").read_text(encoding="utf-8")
    html = html.replace("__DATA_JSON__", json.dumps(payload, separators=(",", ":"), allow_nan=False))
    html = html.replace("__TITLE__", Path(path).stem)
    out.write_text(html, encoding="utf-8")
    print(f"Wrote {out}  ({out.stat().st_size/1e6:.1f} MB)")
    for s in sheets:
        print(f"  {s['name']:<34} {s['nrows']:>6} rows x {len(s['cols']):>3} cols")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("-n", "--sheets", type=int, default=7, help="leading sheets to use (default 7)")
    ap.add_argument("-o", "--out", default=str(HERE / "dashboard.html"))
    a = ap.parse_args()
    build(a.xlsx, a.sheets, Path(a.out))
